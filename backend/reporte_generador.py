"""Genera el .docx del 'Reporte de Trabajo' de los técnicos de EMUNAH.

Encabezado con fecha, técnico, lugar de la operación y responsable del área,
seguido de las entradas que el técnico fue agregando (texto + foto opcional).
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageOps

LOGO_PATH = Path(__file__).resolve().parent / "assets" / "logo_emunah.png"
AZUL = RGBColor(0x1E, 0x4D, 0x8B)
FOTO_MAX_DIM = 1000
FOTO_JPEG_QUALITY = 78


class GeneradorError(Exception):
    pass


def _preparar_foto(src: Path, dest: Path):
    im = Image.open(src)
    im = ImageOps.exif_transpose(im).convert("RGB")
    im.thumbnail((FOTO_MAX_DIM, FOTO_MAX_DIM), Image.LANCZOS)
    im.save(dest, "JPEG", quality=FOTO_JPEG_QUALITY, optimize=True)


def generar_reporte(cfg: dict, out_path: Path, tmp_dir: Path):
    """cfg: fecha, tecnico, lugar, responsable_area y 'entradas' = lista de
    dicts con 'texto' y 'foto' (ruta local o None)."""
    entradas = cfg.get("entradas") or []
    if not entradas:
        raise GeneradorError("No se agregó ninguna entrada")

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)

    if LOGO_PATH.exists():
        doc.add_picture(str(LOGO_PATH), height=Cm(1.6))
    t = doc.add_paragraph()
    r = t.add_run("REPORTE DE TRABAJO")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = AZUL
    p = doc.add_paragraph("Soluciones EMUNAH SAS")
    p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    tabla = doc.add_table(rows=4, cols=2)
    tabla.style = "Table Grid"
    datos = [
        ("Fecha", cfg["fecha"]),
        ("Técnico", cfg["tecnico"]),
        ("Lugar de la operación", cfg["lugar"]),
        ("Responsable del área", cfg["responsable_area"]),
    ]
    for fila, (k, v) in zip(tabla.rows, datos):
        fila.cells[0].width = Cm(5)
        fila.cells[1].width = Cm(12)
        fila.cells[0].paragraphs[0].add_run(k).bold = True
        fila.cells[1].paragraphs[0].add_run(v)

    doc.add_paragraph()
    for i, e in enumerate(entradas, start=1):
        h = doc.add_paragraph()
        rh = h.add_run(f"{i}.")
        rh.bold = True
        rh.font.color.rgb = AZUL
        h.add_run("  " + (e.get("texto") or "").strip())
        foto = e.get("foto")
        if foto:
            jpg = tmp_dir / f"foto_{i}.jpg"
            try:
                _preparar_foto(Path(foto), jpg)
            except Exception:
                raise GeneradorError(f"No se pudo leer la foto de la entrada {i}")
            doc.add_picture(str(jpg), width=Cm(11))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
