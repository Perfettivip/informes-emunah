"""Genera el .xlsx de 'Control Ventas' de Soluciones EMUNAH SAS.

Toma el responsable que diligenció el formulario y la lista de ventas línea
por línea (fecha, cliente, cédula, contacto, descripción, valor base, iva,
total, tipo de pago) y arma un libro de Excel con esa información y los
totales.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

AZUL = "1E4D8B"
GRIS = "6B7280"
BORDE = "D7DCE3"

LOGO_PATH = Path(__file__).resolve().parent / "assets" / "logo_emunah.png"

TIPOS_PAGO = ["Contado", "Transferencia", "Tarjeta"]

COLS_VENTAS = [
    "Fecha", "Cliente", "Cédula", "Contacto", "Descripción",
    "Valor base", "IVA", "Total", "Tipo de pago",
]


class GeneradorError(Exception):
    pass


def _fmt_moneda(cell):
    cell.number_format = '"$" #,##0'


def _borde_fino():
    lado = Side(style="thin", color=BORDE)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def generar_ventas(cfg: dict, out_path: Path):
    """cfg: dict con 'responsable' y cfg['ventas'] = lista de dicts con las
    llaves fecha, cliente, cedula, contacto, descripcion, valor_base, iva,
    total, tipo_pago."""
    ventas = cfg.get("ventas") or []
    if not ventas:
        raise GeneradorError("No se agregó ninguna venta")

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    anchos = [14, 24, 16, 20, 34, 14, 12, 14, 16]
    for col, ancho in zip("ABCDEFGHI", anchos):
        ws.column_dimensions[col].width = ancho

    fila = 1
    if LOGO_PATH.exists():
        img = XLImage(str(LOGO_PATH))
        img.height = 50
        img.width = 130
        ws.add_image(img, "A1")
        ws.row_dimensions[1].height = 40
        fila = 2

    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=9)
    c = ws.cell(row=fila, column=1, value="SOLUCIONES EMUNAH SAS")
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=AZUL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[fila].height = 22
    fila += 1

    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=9)
    c = ws.cell(row=fila, column=1, value="Control Ventas")
    c.font = Font(bold=True, size=11, color=AZUL)
    c.alignment = Alignment(horizontal="center")
    fila += 2

    ws.cell(row=fila, column=1, value="Responsable").font = Font(bold=True, color=GRIS, size=9)
    ws.cell(row=fila, column=2, value=cfg.get("responsable", "")).font = Font(size=11)
    ws.cell(row=fila, column=4, value="Fecha de envío").font = Font(bold=True, color=GRIS, size=9)
    ws.cell(row=fila, column=5, value=cfg.get("fecha_envio", "")).font = Font(size=11)
    fila += 2

    fila_encabezado_tabla = fila
    for i, titulo in enumerate(COLS_VENTAS, start=1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center")
        c.border = _borde_fino()
    fila += 1

    total_base = total_iva = total_general = 0.0
    for v in ventas:
        valor_base = float(v.get("valor_base") or 0)
        iva = float(v.get("iva") or 0)
        total = float(v.get("total") or (valor_base + iva))
        total_base += valor_base
        total_iva += iva
        total_general += total
        valores = [
            v.get("fecha", ""), v.get("cliente", ""), v.get("cedula", ""),
            v.get("contacto", ""), v.get("descripcion", ""),
            valor_base, iva, total, v.get("tipo_pago", ""),
        ]
        for col, val in enumerate(valores, start=1):
            c = ws.cell(row=fila, column=col, value=val)
            c.border = _borde_fino()
            if col in (6, 7, 8):
                _fmt_moneda(c)
        fila += 1

    fila += 1
    ws.cell(row=fila, column=5, value="TOTALES").font = Font(bold=True, size=11)
    for col, valor in ((6, total_base), (7, total_iva), (8, total_general)):
        ct = ws.cell(row=fila, column=col, value=valor)
        ct.font = Font(bold=True, size=11)
        _fmt_moneda(ct)

    ws.freeze_panes = f"A{fila_encabezado_tabla + 1}"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
