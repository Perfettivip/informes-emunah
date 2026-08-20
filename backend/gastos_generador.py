"""Genera el .xlsx de 'Relación de Gastos en Carretera' de Soluciones EMUNAH SAS.

Toma los datos generales del viaje (placa, conductor, fechas, origen/destino,
anticipo, etc.) y la lista de gastos línea por línea (tipo, ciudad, tercero,
cédula, teléfono, valor) que el conductor diligenció en el formulario web, y
arma un libro de Excel con esa información, subtotales por tipo de gasto y
el saldo final frente al anticipo.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "1E4D8B"
AZUL_CLARO = "EAF1FB"
GRIS = "6B7280"
BORDE = "D7DCE3"

TIPOS_GASTO = [
    "Combustible",
    "Lavado",
    "Peajes",
    "Flete carretera",
    "Estadía",
    "Parqueadero",
    "Llantas",
    "Otros",
]

LOGO_PATH = Path(__file__).resolve().parent / "assets" / "logo_emunah.png"

COLS_GASTOS = ["Tipo", "Ciudad", "Tercero", "Cédula / NIT", "Teléfono", "Detalle", "Valor"]


class GeneradorError(Exception):
    pass


def _fmt_moneda(cell):
    cell.number_format = '"$" #,##0'


def _borde_fino():
    lado = Side(style="thin", color=BORDE)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def _fila_dato(ws, fila, etiqueta1, valor1, etiqueta2="", valor2=""):
    ws.cell(row=fila, column=1, value=etiqueta1).font = Font(bold=True, color=GRIS, size=9)
    ws.cell(row=fila, column=2, value=valor1).font = Font(size=11)
    if etiqueta2:
        ws.cell(row=fila, column=4, value=etiqueta2).font = Font(bold=True, color=GRIS, size=9)
        ws.cell(row=fila, column=5, value=valor2).font = Font(size=11)


def generar_gastos(cfg: dict, out_path: Path):
    """cfg: dict con los datos generales del viaje y cfg['gastos'] = lista de dicts
    con las llaves tipo, ciudad, tercero, cedula, telefono, detalle, valor."""
    gastos = cfg.get("gastos") or []
    if not gastos:
        raise GeneradorError("No se agregó ningún gasto")

    wb = Workbook()
    ws = wb.active
    ws.title = "Gastos en carretera"
    for col, ancho in zip("ABCDEFG", [20, 22, 16, 22, 16, 24, 16]):
        ws.column_dimensions[col].width = ancho

    fila = 1
    if LOGO_PATH.exists():
        img = XLImage(str(LOGO_PATH))
        img.height = 50
        img.width = 130
        ws.add_image(img, "A1")
        ws.row_dimensions[1].height = 40
        fila = 2

    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=7)
    c = ws.cell(row=fila, column=1, value="SOLUCIONES EMUNAH SAS")
    c.font = Font(bold=True, size=14, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=AZUL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[fila].height = 22
    fila += 1

    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=7)
    c = ws.cell(row=fila, column=1, value="Relación de Gastos en Carretera")
    c.font = Font(bold=True, size=11, color=AZUL)
    c.alignment = Alignment(horizontal="center")
    fila += 2

    _fila_dato(ws, fila, "Placa", cfg["placa"], "Conductor", cfg["conductor"])
    fila += 1
    _fila_dato(ws, fila, "Cédula conductor", cfg.get("cedula_conductor", ""), "Cliente", cfg.get("cliente", ""))
    fila += 1
    _fila_dato(ws, fila, "Fecha inicio", cfg.get("fecha_inicio", ""), "Fecha fin", cfg.get("fecha_fin", ""))
    fila += 1
    _fila_dato(ws, fila, "Origen", cfg.get("origen", ""), "Destino", cfg.get("destino", ""))
    fila += 1
    _fila_dato(ws, fila, "KM inicial", cfg.get("km_inicial", ""), "KM final", cfg.get("km_final", ""))
    fila += 1
    _fila_dato(ws, fila, "Manifiesto No.", cfg.get("manifiesto", ""), "Guía No.", cfg.get("guia", ""))
    fila += 1
    _fila_dato(ws, fila, "Producto", cfg.get("producto", ""), "Barriles", cfg.get("barriles", ""))
    fila += 1
    ws.cell(row=fila, column=1, value="Anticipo").font = Font(bold=True, color=GRIS, size=9)
    celda_anticipo = ws.cell(row=fila, column=2, value=cfg.get("anticipo", 0))
    _fmt_moneda(celda_anticipo)
    fila += 2

    fila_encabezado_tabla = fila
    for i, titulo in enumerate(COLS_GASTOS, start=1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center")
        c.border = _borde_fino()
    fila += 1

    subtotales = {t: 0.0 for t in TIPOS_GASTO}
    total_gastos = 0.0
    for g in gastos:
        tipo = g.get("tipo") or "Otros"
        valor = float(g.get("valor") or 0)
        subtotales[tipo] = subtotales.get(tipo, 0.0) + valor
        total_gastos += valor
        valores = [tipo, g.get("ciudad", ""), g.get("tercero", ""), g.get("cedula", ""),
                   g.get("telefono", ""), g.get("detalle", ""), valor]
        for col, val in enumerate(valores, start=1):
            c = ws.cell(row=fila, column=col, value=val)
            c.border = _borde_fino()
            if col == 7:
                _fmt_moneda(c)
        fila += 1

    fila += 1
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
    c = ws.cell(row=fila, column=1, value="Subtotales por tipo de gasto")
    c.font = Font(bold=True, color=AZUL, size=10)
    fila += 1
    for tipo in TIPOS_GASTO:
        monto = subtotales.get(tipo, 0.0)
        if monto == 0:
            continue
        ws.cell(row=fila, column=1, value=tipo).font = Font(size=10)
        cm = ws.cell(row=fila, column=2, value=monto)
        _fmt_moneda(cm)
        fila += 1

    fila += 1
    ws.cell(row=fila, column=1, value="TOTAL RELACIÓN DE GASTOS").font = Font(bold=True, size=11)
    ct = ws.cell(row=fila, column=2, value=total_gastos)
    ct.font = Font(bold=True, size=11)
    _fmt_moneda(ct)
    fila += 1

    anticipo = float(cfg.get("anticipo") or 0)
    saldo = total_gastos - anticipo
    fila += 1
    if abs(saldo) < 1:
        ws.cell(row=fila, column=1, value="SALDO").font = Font(bold=True)
        ws.cell(row=fila, column=2, value="Sin saldo pendiente")
    elif saldo > 0:
        ws.cell(row=fila, column=1, value="SALDO A FAVOR DEL CONDUCTOR").font = Font(bold=True, color="1F8A4C")
        cs = ws.cell(row=fila, column=2, value=saldo)
        cs.font = Font(bold=True, color="1F8A4C")
        _fmt_moneda(cs)
    else:
        ws.cell(row=fila, column=1, value="SALDO A FAVOR DE LA EMPRESA").font = Font(bold=True, color="C0392B")
        cs = ws.cell(row=fila, column=2, value=abs(saldo))
        cs.font = Font(bold=True, color="C0392B")
        _fmt_moneda(cs)

    ws.freeze_panes = f"A{fila_encabezado_tabla + 1}"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
